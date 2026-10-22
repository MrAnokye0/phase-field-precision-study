"""make_report.py -- builds the Word report from the tables and figures produced by the study scripts.

Run from the precision-study folder:  python study/make_report.py
Needs python-docx (pip install python-docx). Output: study/report/precision_report.docx
"""
import os
import sys

import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.shared import Cm, Inches, Pt
except ImportError:
    sys.exit("python-docx is missing: run  pip install python-docx")

HERE = os.path.dirname(os.path.abspath(__file__))
TAB = os.path.join(HERE, "tables")
FIG = os.path.join(HERE, "figures")
OUT = os.path.join(HERE, "report")
os.makedirs(OUT, exist_ok=True)
COMMIT = "c90978331a1f83dbf056e70e624dbe098d667d9a"


def rd(name):
    path = os.path.join(TAB, name)
    if not os.path.exists(path):
        sys.exit("missing table: " + path + "  (run the corresponding analysis script first)")
    return pd.read_csv(path)


def fmt(x, nd=3):
    if isinstance(x, str):
        return x
    if x is None or (isinstance(x, (float, np.floating)) and np.isnan(x)):
        return "-"
    if isinstance(x, (int, np.integer, bool, np.bool_)):
        return str(int(x))
    ax = abs(float(x))
    if ax != 0 and (ax < 1e-3 or ax >= 1e5):
        return f"{x:.{nd - 1}e}"
    return f"{x:.{nd}g}"


def pct(x):
    return "-" if (x is None or (isinstance(x, (float, np.floating)) and np.isnan(x))) else fmt(100 * x) + "%"


doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
for s in doc.sections:
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.8)
count = {"t": 0, "f": 0}


def H(text, level=1):
    doc.add_heading(text, level=level)


def P(text, italic=False, size=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    return para


def BL(lead, text):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(lead).bold = True
    para.add_run(text)


def table(headers, rows, caption, font=8):
    count["t"] += 1
    cap = doc.add_paragraph()
    r = cap.add_run(f"Table {count['t']}. {caption}")
    r.bold = True
    r.font.size = Pt(9)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        rr = c.paragraphs[0].add_run(h)
        rr.bold = True
        rr.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
    doc.add_paragraph()


def figure(fname, caption, width=6.2):
    path = os.path.join(FIG, fname)
    if not os.path.exists(path):
        print("WARNING: figure missing, skipped:", fname)
        return
    count["f"] += 1
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    cap = doc.add_paragraph()
    r = cap.add_run(f"Figure {count['f']}. {caption}")
    r.italic = True
    r.font.size = Pt(9)


# ------------------------------------------------------------------ title
doc.add_heading("Numerical precision and reproducibility in a JAX phase-field code", 0)
P("Screening task, CAMMP research positions  |  Edward Gyasi Anokye  |  15 October 2026", italic=True)

# ------------------------------------------------------------------ 1 summary
H("1. Summary")
P("I compared float32 and float64 runs of the open-source JAX phase-field package (Cahn-Hilliard and Allen-Cahn solvers, "
  "ajvetturini/phase-field-models) on its Landau example, with two integrators, and on its simple Wertheim model. "
  "Float64 is the reference throughout. Main results:")
BL("Float32 does not reproduce float64 for the Landau problem with explicit Euler at the shipped step (dt = 0.001). ",
   "At t = 500 the relative L2 field difference is 0.43, 7.4% of pixels are in the opposite phase and the energy differs by 1.9%. "
   "Float64 runs with random initial perturbations from 1e-14 to 1e-7 stay within 1.7e-4 of the reference, so the divergence "
   "comes from float32 arithmetic and not from sensitivity of the system to any small change.")
BL("Smaller time steps make float32 worse. ",
   "At t = 200 the float32 field error falls from 1.9e-2 (dt = 0.001) to 2.0e-4 (dt = 0.01); the log-log slope against dt is "
   "between -1.6 and -2.2. Early in the run an Euler increment is only about 3.5 float32 spacings at dt = 0.001 (31 at dt = 0.01), "
   "and 7.7% of pixels do not change in a step at dt = 0.001.")
BL("The stability limit is the same in both precisions, the timing of failure is not. ",
   "Explicit Euler is stable up to dt = 0.015 and fails at 1/64 = 0.015625 and above in both precisions; float32 fails earlier "
   "(t = 142 against 154 at dt = 0.015625). The code's printed stability term understates the danger.")
BL("Float32 loses conservation and monotonicity. ",
   "The mean density drifts by up to 1.3e-3 relative over t = 500 (float64: 3e-14), the energy rises between snapshots 3 times "
   "(never in float64), and the energy computed inside a float32 run is accurate only to about 4e-7, so energies should be "
   "re-evaluated in float64.")
BL("The semi-implicit spectral integrator (dt = 0.01) stays close to float64 ",
   "(field difference at most 1.1e-2, final energy within 0.07% at t = 500), but its float32 error is larger than Euler's at the "
   "same dt and the mean density still drifts by 2.6e-4 relative.")
BL("Wertheim model (N = 512, first t = 0.2 only): no stalling. ",
   "Float32 differs from float64 in the pattern by 2.2% and in the energy by 4.8e-7, which matches the float32 mean-density drift.")
BL("Reproducibility. ",
   "Float32 results were bit-identical on two different machines (Windows laptop, Linux sandbox) with the same library versions; "
   "float64 results differed at the 1e-13 level.")

# ------------------------------------------------------------------ 2 method
H("2. Scope and method")
P("Source of the code. The shared Drive folder (2. PF JAX) contained only the repository README and demonstration images, not the "
  f"source. I therefore used the public GitHub repository at commit {COMMIT}, unmodified. All scripts are in study/; the exact "
  "environment is in requirements.txt and in every run's meta.json (Windows, CPU only, Python 3.12.10, JAX 0.11.2, NumPy 2.2.4).")
P("Set-up. Precision is selected by float_type in the TOML configuration and is a process-wide JAX setting, so every run is a "
  "separate process. All Landau runs use the Cahn-Hilliard model on a 64 x 64 periodic grid (dx = 1, M = 1, k = 1, "
  "interface_scalar = 2, epsilon = 0.9) and start from the example's initial field, so both precisions start from identical data. "
  "The values in that file are exactly representable in float32 (checked bit for bit), so a float64 run started from float32-rounded "
  "data was identical to the reference and is not discussed further.")
table(["Group", "Runs", "Model / integrator", "dt", "Length"],
      [["Euler main", "euler_f32, euler_f64", "Landau CH, explicit Euler", "0.001", "500,000 steps (t = 500)"],
       ["Controls", "euler_f64_p1e-14, _p1e-10, _p1e-7", "same, float64, perturbed start", "0.001", "500,000 steps"],
       ["dt sweep", "sweep_dt*_f32 / _f64 (11 values)", "explicit Euler", "0.001 to 0.02", "to t = 200"],
       ["Spectral", "spec_f32, spec_f64", "semi-implicit spectral", "0.01", "50,000 steps (t = 500)"],
       ["Wertheim", "wert_f32, wert_f64", "simple Wertheim CH, Euler, N = 512", "1e-5", "20,000 steps (t = 0.2)"]],
      "Experiments.")
P("Metrics. Differences are measured against float64: relative L2 norm of the field difference, maximum pointwise difference, "
  "fraction of pixels with opposite sign (opposite phase) and relative energy difference. Energies and the mean density are "
  "re-evaluated in float64 from the stored fields (an independent NumPy implementation for Landau, which agrees with the package "
  "to 1.5e-16 in float64; the package's own functional in a float64 process for Wertheim). Two controls give context: float64 "
  "runs with random relative perturbations of the initial field (1e-14, 1e-10, 1e-7) measure the sensitivity of the system itself, "
  "and a one-step comparison of the real float32 step with the exact float64 step from the same state measures how well "
  "float32 represents an increment.")

# ------------------------------------------------------------------ 3 results
H("3. Results")
H("3.1 Landau, explicit Euler (dt = 0.001, t = 500)", 2)
P("The Landau example is a spinodal decomposition: a nearly uniform field with 1e-4 noise is linearly unstable (growth rate "
  "estimated at about 0.1 from the linearised equation), the pattern forms around t = 100 and then coarsens.")
tp = rd("euler_timepoints.csv")
rows = []
for _, r in tp[tp["run"] == "euler_f32"].iterrows():
    rows.append(["float32", fmt(r["t"]), fmt(r["rel_L2"]), fmt(r["max_abs_diff"]), pct(r["sign_mismatch"]), fmt(r["dE_rel"])])
for _, r in tp[(tp["t"] == 500) & tp["run"].str.contains("_p1e")].iterrows():
    rows.append(["float64, perturbation " + r["run"].split("_p")[-1], fmt(r["t"]), fmt(r["rel_L2"]), fmt(r["max_abs_diff"]),
                 pct(r["sign_mismatch"]), fmt(r["dE_rel"])])
table(["Run", "t", "rel. L2 difference", "max |difference|", "opposite phase", "rel. energy difference"], rows,
      "Difference to the float64 reference, Euler, dt = 0.001.")
P("Differences are large from the start: 1.9e-4 at t = 20, about 1.8e3 times the difference produced by a 1e-7 initial "
  "perturbation (1.1e-7). The difference first exceeds 1% at t = 58 and 10% at t = 292. It peaks during pattern formation "
  "(t = 100: field 6.1e-2, energy 11.8% lower), recovers partly (about 2% at t = 150 to 200) and rises again during coarsening "
  "(0.11 at t = 300, 0.43 at t = 500). The final energy of the float32 run is -0.08597 against -0.08765 for float64.")
P("The perturbed float64 runs show that this is not generic sensitivity: perturbations of 1e-14, 1e-10 and 1e-7 end at 1.8e-11, "
  "1.7e-7 and 1.7e-4, a constant amplification of about 1.7e3, and no pixel changes phase. A float32 rounding error is of order "
  "1e-7, which on this scale would produce about 1.7e-4 at t = 500, yet float32 differs by 0.43. Float32 arithmetic is therefore "
  "far more disruptive than a one-off perturbation of the same size (Section 3.4 looks at why), and the cause is not the initial data.")
figure("fig1_field_divergence.png", "Relative L2 difference to the float64 reference over time: float32 and the perturbed float64 runs.")
figure("fig4_final_fields.png", "Final fields at t = 500 (top) and their differences to the float64 reference (bottom).")

H("3.2 Energy and mass", 2)
sm = rd("euler_summary.csv").set_index("run")
a, b = sm.loc["euler_f64"], sm.loc["euler_f32"]
table(["Quantity", "float64", "float32"],
      [["Throughput (steps/s)", fmt(a["steps_per_s"]), fmt(b["steps_per_s"])],
       ["Final energy (float64 re-evaluation)", fmt(a["final_E_ref64"]), fmt(b["final_E_ref64"])],
       ["Final change of the mean density", fmt(a["final_mass_drift"]), fmt(b["final_mass_drift"])],
       ["Largest |change| of the mean density", fmt(a["max_abs_mass_drift"]), fmt(b["max_abs_mass_drift"])],
       ["Energy increases between snapshots", fmt(a["E_increases_ref64"]), fmt(b["E_increases_ref64"])],
       ["Largest relative error of the in-run energy", fmt(a["max_E_readout_rel_err"]), fmt(b["max_E_readout_rel_err"])]],
      "Energy and mass diagnostics, Euler, dt = 0.001, t = 500 (mean density 0.01).")
P("Float32 does not conserve the mean density: it changes by up to 1.3e-5 (1.3e-3 relative) while float64 stays within 3e-16. "
  "The float32 energy rose between snapshots 3 times, which a gradient flow should not do, and the energy computed inside the "
  "float32 run is accurate to only about 4e-7. Because the energies in this problem start at about -4.5e-5, the package's own text "
  "output (energy.dat, 8 decimals) keeps about four significant digits; I therefore stored full-precision values and re-evaluated "
  "all energies in float64.")
figure("fig2_energy.png", "Energy over time (top) and relative energy difference to the float64 reference (bottom).")

H("3.3 Stability: time-step sweep", 2)
st = rd("sweep_stability.csv")


def stab(x):
    return f"stable to t = {fmt(x['t_last_ok'])}" if x["status"] == "ok" else f"blow-up, first bad snapshot t = {fmt(x['t_first_bad'])}"


rows = []
for dt in sorted(st["dt"].unique()):
    a = st[(st["dt"] == dt) & (st["prec"] == "f32")].iloc[0]
    b = st[(st["dt"] == dt) & (st["prec"] == "f64")].iloc[0]
    rows.append([f"{dt:g}", stab(a), stab(b), fmt(a["E_increases"]), fmt(b["E_increases"])])
table(["dt", "float32", "float64", "energy increases (f32)", "energy increases (f64)"], rows,
      "Stability of explicit Euler against dt (run to t = 200).")
P("Both precisions are stable up to dt = 0.015 and fail at dt = 0.015625 (= 1/64), 0.016 and 0.02, the last within the first two "
  "time units. Float32 fails earlier than float64 at 0.015625 (first bad snapshot at t = 142 against 154) and at 0.016 (128 against 134); "
  "my explanation, not tested, is that the larger rounding noise of float32 seeds the unstable mode sooner. For dt up to 0.015 the "
  "grid-scale (checkerboard) amplitude never exceeded its initial value of 7.5e-7, so runs that stayed finite contain no hidden "
  "oscillation. The only stable run with energy increases is float32 at dt = 0.001.")
P("The observed limit lies between 0.015 and 0.015625, close to 2/(64 M kappa) = 1/64, the bound from the interface term alone "
  "(kappa = interface_scalar x k = 2). The code prints a \"stability term\" (2 M k dt / dx^4) that should be \"much smaller than 1\": "
  "it is 0.04 at dt = 0.02, a step that diverges within 100 steps, and 0.031 at the limit, so it understates the danger by more "
  "than an order of magnitude. A hand estimate that included the bulk term predicted failure already at dt = 0.015; this was not "
  "observed, so I do not claim to understand the bulk contribution.")

H("3.4 Why smaller steps hurt float32", 2)
sp = rd("sweep_precision.csv")
rows = [[f"{r['dt']:g}", fmt(r["relL2_t20"]), fmt(r["relL2_t50"]), fmt(r["relL2_t200"]), fmt(r["final_dE_rel"]), fmt(r["f32_mass_drift"])]
        for _, r in sp.iterrows()]
table(["dt", "rel. L2, t = 20", "t = 50", "t = 200", "rel. energy difference at t = 200", "float32 mean-density change"], rows,
      "Float32 against float64 at the same dt (a dash means the run had stopped).")
P("The float32 error falls steeply as dt grows: at t = 200 it is 1.9e-2 at dt = 0.001, 3.8e-3 at 0.0025, 1.8e-3 at 0.005 and 2.0e-4 at 0.01. "
  "The log-log slope of the t = 20 error against dt is -1.6 (dt up to 0.01, four points) to -2.05 (seven points, including steps near the "
  "stability limit); at t = 50 it is -1.7 to -2.2. The float32 mean-density drift behaves the same way (-7.0e-6 at dt = 0.001, "
  "-1.5e-7 at dt = 0.01).")
sg = rd("stagnation.csv")
rows = [[f"{r['dt']:g}", fmt(r["t"]), fmt(r["median_incr_over_ulp"]), pct(r["frac_stagnated"]),
         fmt(r["share_of_movement_lost"]), fmt(r["rel_step_error"])]
        for _, r in sg[sg["t"].isin([20, 50, 100, 200])].sort_values(["dt", "t"]).iterrows()]
table(["dt", "t", "median increment / float32 spacing", "pixels unchanged in a step", "share of movement lost", "relative step error"], rows,
      "One real float32 step compared with the exact float64 step from the same state.")
P("Table 6 shows why. Early in the run (t = 20) the increment added in a step is small compared with the float32 spacing near 0.01: "
  "a median of 3.5 spacings at dt = 0.001 against 31 at dt = 0.01. The rounding error of a step relative to its increment is 9.3% "
  "against 0.69%, and 7.7% against 0.98% of pixels do not change at all (the increment is below half a spacing and is lost). "
  "By t = 100 nothing stagnates; stagnation returns weakly in the coarsening stage, but those pixels carry less than 1e-4 of the total "
  "movement. The per-step error scales roughly as dt^-1.1. This is consistent with the observed dependence if errors accumulate "
  "coherently over the number of steps, but I did not test that. The one-step results matched an independent run on a second "
  "machine to the printed digits.")
figure("fig6_error_vs_dt.png", "Float32-against-float64 field difference at t = 20 and t = 50 against dt (log-log).")

H("3.5 Semi-implicit spectral integrator (dt = 0.01, t = 500)", 2)
ptp = rd("pair_spectral_dt0.01_timepoints.csv")
rows = [[fmt(r["t"]), fmt(r["rel_L2"]), fmt(r["max_abs_diff"]), pct(r["sign_mismatch"]), fmt(r["dE_rel"])] for _, r in ptp.iterrows()]
table(["t", "rel. L2 difference", "max |difference|", "opposite phase", "rel. energy difference"], rows,
      "Spectral integrator, float32 against float64.")
ps = rd("pair_spectral_dt0.01_summary.csv").set_index("run")
a, b = ps.loc["spec_f64"], ps.loc["spec_f32"]
table(["Quantity", "float64", "float32"],
      [["Throughput (steps/s)", fmt(a["steps_per_s"]), fmt(b["steps_per_s"])],
       ["Final energy", fmt(a["final_E"]), fmt(b["final_E"])],
       ["Final change of the mean density", fmt(a["final_mass_drift"]), fmt(b["final_mass_drift"])],
       ["Energy increases between snapshots", fmt(a["E_increases"]), fmt(b["E_increases"])],
       ["Largest relative error of the in-run energy", fmt(a["max_E_readout_rel_err"]), fmt(b["max_E_readout_rel_err"])]],
      "Spectral integrator, summary.")
P("Float32 stays close to float64 here: the field difference is at most 1.1e-2 (t = 300) and 6.3e-3 at t = 500, at most 0.27% of "
  "pixels are in the opposite phase and the final energies differ by 0.07%. The energy decreases monotonically in both precisions. "
  "The mean density still drifts by -2.6e-6 (2.6e-4 relative) in float32 against about 2e-16 in float64, and float32 is 1.27 times faster. "
  "At the same dt = 0.01 the Euler float32 error at t = 200 is 2.0e-4 against 7.0e-3 for the spectral scheme, so the spectral scheme is "
  "less accurate in float32; my untested explanation is that the FFT's rounding affects the whole field every step and not only the increment.")
figure("pair_spectral_dt0.01.png", "Spectral integrator: field difference, energy and mean-density drift for float32 against float64.")

H("3.6 Simple Wertheim model (N = 512, dt = 1e-5, t = 0.2)", 2)
wc = rd("wertheim_euler_comparison.csv")
rows = [[fmt(r["t"]), fmt(r["relL2_field"]), fmt(r["relL2_deviation"]), fmt(r["dE_rel"]), fmt(r["mass_drift_rel_test"]),
         fmt(r["disp_ratio"]), fmt(r["Edrop_ratio"]), pct(r["frozen_test"])] for _, r in wc.iterrows()]
table(["t", "rel. L2 (field)", "rel. L2 (deviation from mean)", "rel. energy difference", "float32 mean-density drift (rel.)",
       "displacement ratio f32/f64", "energy-drop ratio f32/f64", "float32 pixels frozen"], rows,
      "Wertheim, float32 against float64 (snapshots 2000 steps apart).")
P("The shipped Wertheim example runs 1e8 steps on a GPU; I ran 20,000 (t = 0.2; 130 s in float32 and 432 s in float64, so float32 is 3.3 times "
  "faster at this size). Float32 does not stall in this window: the displacement from the starting state matches float64 to within "
  "1.2% (ratio 0.9998 to 1.012) with correlation at least 0.9999. A small fraction of float32 pixels is frozen over each 2000-step "
  "interval (up to 0.27%, falling to 0.11% by t = 0.2) against none in float64. The pattern (deviation from the mean) differs by 2.2% "
  "at t = 0.2, growing roughly linearly in time. The float32 mean density drifts by 4.8e-7 relative, accelerating, and the "
  "energy difference (4.8e-7) matches this drift to within about 1% at every snapshot, so the float32 energy error here is almost entirely "
  "the mass drift; I did not separate the two further. The float32 energy decreases 10.6% more than the float64 energy by t = 0.2. "
  "The energy never increases in either run, and the in-run float32 energy has a relative error of 6.9e-8. The README's remark that float32 stalls "
  "on this model concerns the long-time regime, which this run does not reach. The Wertheim start is not exactly representable in float32 "
  "(relative L2 2.1e-8, or 3.6e-6 for the deviation from the mean) and I did not run the float64-from-rounded-start control for this model.")
figure("wertheim_euler.png", "Wertheim: energy change, field differences, dynamics comparison and mean-density drift.")

H("3.7 Reproducibility across machines", 2)
P("I repeated the spectral runs on a second machine (a Linux sandbox, same JAX 0.11.2 and NumPy 2.2.4, CPU). The float32 result file was "
  "bit-identical (same SHA-256 hash). The float64 results differed in the last bits: the sums of the final field differ by 3.8e-12 "
  "(9e-14 relative), eleven orders of magnitude below the float32-against-float64 difference (7e-3). Other printed results from the same "
  "scripts (Landau smoke runs, Wertheim speed test, one-step diagnostic) also matched to the printed digits. This is one machine pair "
  "with identical library versions and no GPU; it should not be read as a general guarantee.")

# ------------------------------------------------------------------ 4 impact
H("4. Impact of numerical precision and recommendations")
P("What precision changed: pattern-level agreement for the Landau problem at small dt, conservation of the mean density, the energy "
  "diagnostics (monotonicity and accuracy of the printed energy), the time at which an unstable step size fails, and speed "
  "(float32 is 1.2 to 1.3 times faster at 64 x 64 and 3.3 times faster at 512 x 512). What it did not change: the initial data "
  "(exactly representable for Landau), the stability limit, and the qualitative dynamics in the early Wertheim run.")
P("Recommendations (the first three follow directly from the measurements; the last two are suggestions I did not test):")
BL("Use float64 ", "when the energy, mass conservation or the detailed pattern at long times matter. Use float32 only for qualitative "
   "work, and then with the largest time step that is stable: at t = 200 dt = 0.01 reduced the float32 field error about 100-fold relative to dt = 0.001.")
BL("Re-evaluate energies and means in float64 ", "from float32 states (or use compensated sums); the energy printed by the float32 run is limited to about 1e-7.")
BL("Always monitor the mean density ", "in float32 runs: the drift is easy to detect and was present in every float32 run.")
BL("Consider re-normalising the mean density ", "periodically in float32 runs (not tested).")
BL("Record seeds, versions and hardware ", "with every run and fix the initial data, as done here.")

H("5. Observations on the code", 2)
table(["Observation", "Status"],
      [["The printed Euler stability term understates the instability (it is 0.04 at dt = 0.02, a step that diverges within 100 steps).", "verified"],
       ["FreeEnergyModel tests hasattr(config, 'float_type') on a plain dict, which is always False, so its floor/regularisation "
        "constant stays a float32 1e-7 even for float64 runs (used by Wertheim, not by Landau).", "verified by running; effect on Wertheim results not tested"],
       ["The CPU path writes energy.dat with 8 decimals (energy) and 5 (mass): too coarse to compare precisions on the Landau example.", "verified"],
       ["The JAX/accelerator path stores trajectory frames as float32 whatever float_type is.", "from reading the code, not tested"],
       ["Without a seed in the configuration a random seed is drawn per run (matters only when the initial field is generated, not loaded); "
        "SimulationManagerNoWrite takes its seed from the number of steps.", "from reading the code, not tested"]],
      "Observations on the code.")

H("6. Limitations")
for t in ["Landau results come from one 64 x 64 grid and one initial condition; the dt sweep and spectral runs are single runs.",
          "Wertheim covers only the first t = 0.2 of the example's 1e8 steps, so the long-time stalling is not tested.",
          "CPU only; no GPU runs. Bit-level reproducibility was checked on one machine pair with the same library versions.",
          "The mechanism behind the dt dependence (increment quantisation accumulating coherently) is consistent with the data but not proven, "
          "and the stability estimate including the bulk term disagrees with the observations.",
          "The run times are single measurements on a laptop."]:
    doc.add_paragraph(t, style="List Bullet")

# ------------------------------------------------------------------ appendix
H("Appendix: additional figures and reproduction")
figure("fig3_mass_drift.png", "Mean-density drift for the float32, float64 and perturbed float64 Euler runs.")
figure("fig5_sweep_divergence.png", "Float32 against float64 field difference over time for each dt in the sweep.")
P("Reproduction. With the environment in requirements.txt, running .\\study\\run_all.ps1 from the folder containing study/ and "
  "phase-field-models/ repeats every run and rebuilds all tables and figures (about 20 to 30 minutes on a laptop CPU); "
  "python study\\make_report.py rebuilds this document. Each run's folder holds its configuration (config.toml), metadata (meta.json), "
  "log (log.csv) and snapshots (fields.npy).")

path = os.path.join(OUT, "precision_report.docx")
doc.save(path)
print("Wrote", path, f"({count['t']} tables, {count['f']} figures)")
